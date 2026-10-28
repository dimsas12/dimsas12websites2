"""Excel report for one game day: picks (with live stake formulas and a results
tracker), game predictions vs. market, the full odds board, injuries and
player prop projections.

Prices come from the most recent `elbet picks` run (reports/markets_*.csv),
so run `elbet picks` first for fresh odds.
"""

from __future__ import annotations

import glob
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from . import config, dataset, injuries, pipeline
from .config import SETTINGS
from .models import GameModel
from .props import PROP_STATS

FONT = "Arial"
HEAD_FILL = PatternFill("solid", fgColor="1F3864")
INPUT_FILL = PatternFill("solid", fgColor="FFFF00")
PICK_FILL = PatternFill("solid", fgColor="E2EFDA")
BLUE = "0000FF"
GREEN = "008000"
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)

MARKET_LABEL = {"moneyline": "Moneyline", "spread": "Spread", "total": "Total"}


def _f(bold=False, color="000000", size=10, italic=False):
    return Font(name=FONT, bold=bold, color=color, size=size, italic=italic)


def _header(ws, row: int, headers: list[str], widths: list[int] | None = None) -> None:
    for j, h in enumerate(headers, 1):
        c = ws.cell(row=row, column=j, value=h)
        c.font = _f(bold=True, color="FFFFFF")
        c.fill = HEAD_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BOX
    ws.row_dimensions[row].height = 30
    if widths:
        for j, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(j)].width = w


def _cell(ws, row, col, value, fmt=None, color="000000", bold=False, fill=None):
    c = ws.cell(row=row, column=col, value=value)
    c.font = _f(bold=bold, color=color)
    c.border = BOX
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = fill
    return c


def _title(ws, text: str, sub: str) -> None:
    ws["A1"] = text
    ws["A1"].font = _f(bold=True, size=14)
    ws["A2"] = sub
    ws["A2"].font = _f(italic=True, color="595959")


def _bet_label(r) -> str:
    if r.market == "total":
        return f"{r.selection.capitalize()} {r.line:g}"
    if r.market == "spread":
        return f"{r.pick} {r.line:+g}"
    return f"{r.pick} to win"


def _latest_board(game_ids: set[str]) -> tuple[pd.DataFrame, str | None]:
    files = sorted(glob.glob(str(config.REPORTS_DIR / "markets_*.csv")))
    for f in reversed(files):
        b = pd.read_csv(f)
        b = b[b["game_id"].isin(game_ids)]
        if not b.empty:
            stamp = Path(f).stem.replace("markets_", "")
            return b, stamp
    return pd.DataFrame(), None


def _picks_from_board(board: pd.DataFrame) -> pd.DataFrame:
    bet = SETTINGS.bet
    p = board[(board["ev"] >= bet.min_edge) & (board["model_p"] >= bet.min_prob)
              & (board["odds"] <= bet.max_odds)].sort_values("ev", ascending=False)
    return p.drop_duplicates("game") if bet.one_bet_per_game else p


def export(day: pd.Timestamp, out: Path | None = None, fetch_injuries: bool = True) -> Path:
    feats = pipeline.load_features()
    model = GameModel.load(pipeline.GAME_MODEL_PATH)
    day = pd.Timestamp(day).normalize()
    up = feats[(feats["date"].dt.normalize() == day)].sort_values("date")
    if up.empty:
        raise ValueError(f"no games on {day:%Y-%m-%d}")
    inj = injuries.load_injuries(fetch=fetch_injuries)
    up, absent = pipeline.apply_absences(up, inj)
    pred = pipeline.predict_games(up, model)
    coef = model.margin_model.raw_coefficients(model.margin_features)["avail_diff"]
    tcoef = model.total_model.raw_coefficients(model.total_features)["missing_pts_total"]
    board, stamp = _latest_board(set(pred["game_id"]))
    picks = _picks_from_board(board) if not board.empty else board

    wb = Workbook()
    _sheet_picks(wb.active, day, picks, stamp)
    _sheet_games(wb.create_sheet("Games"), pred, board, coef, tcoef)
    _sheet_board(wb.create_sheet("Odds board"), board, stamp)
    _sheet_injuries(wb.create_sheet("Injuries"), absent, inj, pred, coef, tcoef)
    props_model = pipeline.load_props_model()
    allprops = _all_props(pred, inj, props_model)
    _sheet_props(wb.create_sheet("Player props"), pred, allprops)
    n_ladder = _sheet_ladder(wb.create_sheet("Prop ladder"), allprops, props_model)
    _sheet_checker(wb.create_sheet("Prop checker"), allprops, n_ladder)
    _sheet_notes(wb.create_sheet("Notes"), model, stamp, inj)

    out = out or config.REPORTS_DIR / f"euroleague_{day:%Y-%m-%d}.xlsx"
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out


# ---------------------------------------------------------------- sheets

def _sheet_picks(ws, day, picks: pd.DataFrame, stamp) -> None:
    ws.title = "Picks"
    _title(ws, f"EuroLeague picks - {day:%A %d %B %Y}",
           f"Odds snapshot {stamp or 'n/a'} UTC. Paper trading: the model has no proven edge yet.")
    # Inputs
    ws["A4"], ws["A5"], ws["A6"] = "Bankroll (EUR)", "Kelly fraction", "Max stake per bet"
    for r, v, fmt in ((4, 1000, "#,##0"), (5, SETTINGS.bet.kelly_fraction, "0.00"),
                      (6, SETTINGS.bet.max_stake_pct, "0.0%")):
        ws.cell(row=r, column=1).font = _f(bold=True)
        c = _cell(ws, r, 2, v, fmt, color=BLUE, fill=INPUT_FILL)
    ws["C4"] = "<- yellow cells are inputs; stakes, EV and profit recalculate"
    ws["C4"].font = _f(italic=True, color="595959")
    ws["B5"].comment = Comment("0.25 = quarter Kelly (default in elbet).", "elbet")

    headers = ["Time (UTC)", "Game", "Market", "Bet", "Odds (net)", "Book", "Model win %",
               "Market win % (no-vig)", "Push %", "Fair odds", "EV per 1 staked",
               "Full Kelly %", "Stake (EUR)", "Closing odds", "Result (W/L/P)",
               "Profit (EUR)", "Price CLV"]
    widths = [11, 11, 11, 15, 8, 14, 10, 12, 8, 9, 10, 10, 10, 10, 10, 11, 10]
    _header(ws, 8, headers, widths)
    dv = DataValidation(type="list", formula1='"W,L,P"', allow_blank=True)
    ws.add_data_validation(dv)
    r = 9
    picks = picks.sort_values("commence") if not picks.empty else picks
    for p in picks.itertuples(index=False):
        _cell(ws, r, 1, pd.Timestamp(p.commence).strftime("%H:%M"))
        _cell(ws, r, 2, p.game)
        _cell(ws, r, 3, MARKET_LABEL[p.market])
        _cell(ws, r, 4, _bet_label(p), bold=True)
        _cell(ws, r, 5, float(p.odds), "0.00", color=BLUE)
        _cell(ws, r, 6, p.book)
        _cell(ws, r, 7, round(float(p.model_p), 4), "0.0%", color=BLUE)
        _cell(ws, r, 8, None if pd.isna(p.market_p) else round(float(p.market_p), 4), "0.0%",
              color=BLUE)
        _cell(ws, r, 9, round(float(p.p_push), 4), "0.0%", color=BLUE)
        _cell(ws, r, 10, f"=1/G{r}", "0.00")
        _cell(ws, r, 11, f"=G{r}*(E{r}-1)-(1-G{r}-I{r})", "0.0%")
        _cell(ws, r, 12, f"=MAX(0,K{r}/((E{r}-1)*(1-I{r})))", "0.0%")
        _cell(ws, r, 13, f"=ROUND(MIN(L{r}*$B$5,$B$6)*$B$4,2)", "#,##0.00")
        _cell(ws, r, 14, None, "0.00", color=BLUE, fill=INPUT_FILL)
        _cell(ws, r, 15, None, color=BLUE, fill=INPUT_FILL)
        dv.add(f"O{r}")
        _cell(ws, r, 16, f'=IF(O{r}="W",M{r}*(E{r}-1),IF(O{r}="L",-M{r},0))', "#,##0.00")
        _cell(ws, r, 17, f'=IF(N{r}="","",E{r}/N{r}-1)', "0.0%")
        r += 1
    last = r - 1
    if picks.empty:
        ws.cell(row=9, column=1, value="No bets cleared the minimum edge for this day.").font = \
            _f(italic=True)
        return
    ws.cell(row=r, column=12, value="Total").font = _f(bold=True)
    _cell(ws, r, 13, f"=SUM(M9:M{last})", "#,##0.00", bold=True)
    _cell(ws, r, 16, f"=SUM(P9:P{last})", "#,##0.00", bold=True)
    _cell(ws, r, 17, f'=IF(COUNT(Q9:Q{last})=0,"",AVERAGE(Q9:Q{last}))', "0.0%", bold=True)
    ws.cell(row=r + 1, column=12, value="ROI").font = _f(bold=True)
    _cell(ws, r + 1, 16, f'=IF(SUMPRODUCT((O9:O{last}<>"")*M9:M{last})=0,"",'
                         f'P{r}/SUMPRODUCT((O9:O{last}<>"")*M9:M{last}))', "0.0%", bold=True)
    note = r + 3
    for i, t in enumerate([
        "How to use: after each game set Result to W, L or P (push); Profit and ROI update.",
        "Enter the Closing odds (same bet, same line, just before tip-off) to see Price CLV = "
        "your odds / closing odds - 1. Beating the close on average is the best sign of an edge.",
        "Blue = hard-coded model/market numbers from the snapshot; black = formulas.",
        "Odds (net) = price after exchange commission (Betfair 5%, Matchbook 2%); offshore books "
        "(1xBet, GTbets, BetOnline, MyBookie, Marathonbet...) are never used as best price.",
        "One bet per game (best EV). Totals are the model's weakest market: treat those edges "
        "with extra caution.",
    ]):
        ws.cell(row=note + i, column=1, value=t).font = _f(italic=True, color="595959")
    ws.freeze_panes = "A9"


def _sheet_games(ws, pred: pd.DataFrame, board: pd.DataFrame, coef: float,
                 tcoef: float) -> None:
    _title(ws, "Game predictions vs market",
           "Pure model (with injuries), before the 50/50 market blend used on the Picks sheet. "
           "Market = consensus no-vig across bookmakers.")
    headers = ["Time (UTC)", "Home", "Away", "Model home win %", "Fair home odds",
               "Fair away odds", "Pred. margin (home)", "Model spread (home)", "Market spread (home)",
               "Spread diff (pts)", "Pred. total", "Market total", "Total diff (pts)",
               "Pred. score", "Market home win %", "Win % diff", "Injury edge (pts, + = home)",
               "Injury effect on total (pts)"]
    _header(ws, 4, headers, [10, 22, 22, 10, 9, 9, 10, 10, 10, 9, 9, 9, 9, 11, 10, 9, 11])
    r = 5
    for g in pred.itertuples(index=False):
        b = board[board["game_id"] == g.game_id] if not board.empty else board
        sp = b[(b["market"] == "spread") & (b["selection"] == "home")]
        to = b[(b["market"] == "total") & (b["selection"] == "over")]
        ml = b[(b["market"] == "moneyline") & (b["selection"] == "home")]
        _cell(ws, r, 1, g.date.strftime("%H:%M"))
        _cell(ws, r, 2, g.home_name)
        _cell(ws, r, 3, g.away_name)
        _cell(ws, r, 4, round(g.p_home, 4), "0.0%", color=BLUE)
        _cell(ws, r, 5, f"=1/D{r}", "0.00")
        _cell(ws, r, 6, f"=1/(1-D{r})", "0.00")
        _cell(ws, r, 7, round(g.pred_margin, 2), "0.0", color=BLUE)
        _cell(ws, r, 8, f"=ROUND(-G{r}*2,0)/2", "+0.0;-0.0;0.0")
        _cell(ws, r, 9, float(sp["line"].iloc[0]) if not sp.empty else None, "+0.0;-0.0;0.0",
              color=BLUE)
        _cell(ws, r, 10, f'=IF(I{r}="","",I{r}-H{r})', "+0.0;-0.0;0.0")
        _cell(ws, r, 11, round(g.pred_total, 1), "0.0", color=BLUE)
        _cell(ws, r, 12, float(to["line"].iloc[0]) if not to.empty else None, "0.0", color=BLUE)
        _cell(ws, r, 13, f'=IF(L{r}="","",K{r}-L{r})', "+0.0;-0.0;0.0")
        _cell(ws, r, 14, f'=ROUND((K{r}+G{r})/2,0)&" - "&ROUND((K{r}-G{r})/2,0)')
        mp = ml["market_p"].iloc[0] if not ml.empty else None
        _cell(ws, r, 15, None if mp is None or pd.isna(mp) else round(float(mp), 4), "0.0%",
              color=BLUE)
        _cell(ws, r, 16, f'=IF(O{r}="","",D{r}-O{r})', "+0.0%;-0.0%;0.0%")
        _cell(ws, r, 17, round(float(g.avail_diff) * coef, 2), "+0.0;-0.0;0.0", color=BLUE)
        _cell(ws, r, 18, round(float(g.missing_pts_total) * tcoef, 2), "+0.0;-0.0;0.0",
              color=BLUE)
        r += 1
    ws.cell(row=r + 1, column=1, value=(
        "Spread convention: negative = home team favoured (e.g. -6.5 means home must win by 7+). "
        "Injury edge = points of margin the home side gains from absences (away missing "
        "players minus home missing players).")).font = _f(italic=True, color="595959")
    ws.freeze_panes = "D5"


def _sheet_board(ws, board: pd.DataFrame, stamp) -> None:
    _title(ws, "Full odds board", f"Best price across bookmakers, snapshot {stamp or 'n/a'} UTC")
    headers = ["Time (UTC)", "Game", "Market", "Bet", "Best odds (net)", "Book", "Model win %",
               "Market win % (no-vig)", "Push %", "Fair odds", "EV per 1 staked", "Pick?"]
    _header(ws, 4, headers, [10, 11, 11, 15, 9, 14, 10, 12, 8, 9, 11, 8])
    picks = _picks_from_board(board) if not board.empty else board
    keys = set(zip(picks.get("game", []), picks.get("market", []), picks.get("selection", [])))
    r = 5
    for p in board.itertuples(index=False):
        is_pick = (p.game, p.market, p.selection) in keys
        fill = PICK_FILL if is_pick else None
        _cell(ws, r, 1, pd.Timestamp(p.commence).strftime("%H:%M"), fill=fill)
        _cell(ws, r, 2, p.game, fill=fill)
        _cell(ws, r, 3, MARKET_LABEL[p.market], fill=fill)
        _cell(ws, r, 4, _bet_label(p), bold=is_pick, fill=fill)
        _cell(ws, r, 5, float(p.odds), "0.00", color=BLUE, fill=fill)
        _cell(ws, r, 6, p.book, fill=fill)
        _cell(ws, r, 7, round(float(p.model_p), 4), "0.0%", color=BLUE, fill=fill)
        _cell(ws, r, 8, None if pd.isna(p.market_p) else round(float(p.market_p), 4), "0.0%",
              color=BLUE, fill=fill)
        _cell(ws, r, 9, round(float(p.p_push), 4), "0.0%", color=BLUE, fill=fill)
        _cell(ws, r, 10, f"=1/G{r}", "0.00", fill=fill)
        _cell(ws, r, 11, f"=G{r}*(E{r}-1)-(1-G{r}-I{r})", "+0.0%;-0.0%;0.0%", fill=fill)
        _cell(ws, r, 12, "YES" if is_pick else "", bold=is_pick, fill=fill)
        r += 1
    ws.cell(row=r + 1, column=1, value=(
        f"Pick = EV >= {SETTINGS.bet.min_edge:.0%}, win % >= {SETTINGS.bet.min_prob:.0%}, odds <= "
        f"{SETTINGS.bet.max_odds:g}, best EV only per game. Model win % is blended 50/50 with the "
        "market (--market-weight 0.5).")).font = _f(italic=True, color="595959")
    ws.freeze_panes = "A5"


def _sheet_injuries(ws, absent, inj, pred, coef, tcoef) -> None:
    _title(ws, "Injuries applied", "Source: basketballsphere.com EuroLeague injury report "
           "(club reports), plus data/injuries/manual.csv")
    headers = ["Game", "Team", "Player", "Status", "Chance out", "Exp. minutes",
               "Value (model units)", "Margin impact (pts)", "Scoring value (pts)",
               "Total impact (pts)", "Injury", "Expected back"]
    _header(ws, 4, headers, [11, 8, 24, 12, 10, 10, 11, 11, 11, 11, 34, 16])
    ws["J3"] = tcoef
    ws["I3"] = "Total pts per pt:"
    ws["I3"].font = _f(bold=True)
    ws["J3"].font = _f(color=BLUE)
    ws["J3"].number_format = "0.000"
    ws["J3"].comment = Comment("Learned from 15 seasons: game-total points per point of "
                               "missing scoring above replacement (teammates absorb most).",
                               "elbet")
    ws["H3"] = coef
    ws["G3"] = "Points per unit:"
    ws["G3"].font = _f(bold=True)
    ws["H3"].font = _f(color=BLUE)
    ws["H3"].number_format = "0.000"
    ws["H3"].comment = Comment("Learned by the game model from 15 seasons of box scores: margin "
                               "points per unit of missing-player value.", "elbet")
    games = dict(zip(pred["game_id"], pred["home"] + "-" + pred["away"]))
    r = 5
    if absent is None or absent.empty:
        ws.cell(row=5, column=1, value="No rotation players listed as absent.").font = _f()
        return
    feed = {str(p).lower(): row for p, row in zip(inj["player"], inj.itertuples(index=False))}
    from .pipeline import match_player
    for a in absent.itertuples(index=False):
        hit = next((n for n in inj["player"] if match_player(n, [a.player])), None)
        src = feed.get(str(hit).lower()) if hit else None
        _cell(ws, r, 1, games.get(a.game_id, a.game_id))
        _cell(ws, r, 2, a.team)
        _cell(ws, r, 3, a.player)
        _cell(ws, r, 4, src.status if src is not None else "Out (manual)")
        _cell(ws, r, 5, float(a.p_out), "0%", color=BLUE)
        _cell(ws, r, 6, round(float(a.exp_min), 1), "0.0", color=BLUE)
        _cell(ws, r, 7, round(float(a.value), 2), "0.00", color=BLUE)
        _cell(ws, r, 8, f"=E{r}*G{r}*$H$3", "0.00")
        _cell(ws, r, 9, round(float(a.pts_value), 2), "0.00", color=BLUE)
        _cell(ws, r, 10, f"=E{r}*I{r}*$J$3", "+0.00;-0.00;0.00")
        _cell(ws, r, 11, src.injury if src is not None else "")
        _cell(ws, r, 12, src.back if src is not None else "")
        r += 1
    ws.freeze_panes = "A5"


STAT_NAMES = {"pts": "PTS", "reb": "REB", "ast": "AST", "fgm3": "3PM", "pra": "PRA"}


def _all_props(pred: pd.DataFrame, inj: pd.DataFrame, props_model) -> pd.DataFrame:
    _, players = dataset.load_tables()
    out = []
    for _, g in pred.iterrows():
        absences = pipeline.game_absences(g, inj)
        proj = pipeline.game_props(g, players, props_model, absences,
                                   {"home": g["pred_home_pts"], "away": g["pred_away_pts"]})
        if not proj.empty:
            proj["game"] = f"{g['home']}-{g['away']}"
            proj["team_pts"] = np.where(proj["side"] == "home", g["pred_home_pts"],
                                        g["pred_away_pts"])
            out.append(proj)
    return pd.concat(out, ignore_index=True) if out else pd.DataFrame()


def _sheet_props(ws, pred, allprops: pd.DataFrame) -> None:
    _title(ws, "Player prop projections",
           "Projections are IF THE PLAYER PLAYS (how books settle props). Chance plays = recent "
           "appearance rate x (1 - injury probability). SD and P(over) include minutes and "
           "shooting-rate uncertainty. Ruled-out players are listed in grey under each team with "
           "no projection and are not in the totals. Check rows: listed players normally cover ~194-197 of a "
           "team's ~201 minutes, and less with many absences (history: ~180 with 4 out) because "
           "call-ups and returning players without recent games take the rest; their points "
           "are compared with the game model's score (forcing them equal tested worse).")
    stats = ["pts", "reb", "ast", "fgm3", "pra"]
    headers = ["Game", "Team", "Player", "Chance plays", "Minutes if plays", "Injury P(out)"]
    for st in stats:
        n = STAT_NAMES[st]
        headers += [f"{n} mean", f"{n} SD", f"{n} fair line", f"{n} P(over fair)"]
    _header(ws, 4, headers, [10, 6, 24, 8, 9, 8] + [7, 6, 7, 8] * len(stats))
    if allprops.empty:
        return
    r = 5
    grey = _f(color="808080", italic=True)
    for (game, team), d_all in allprops.groupby(["game", "team"], sort=False):
        d = d_all[d_all["p_play"] > 0]
        ruled_out = d_all[d_all["p_play"] <= 0]
        first = r
        for p in d.sort_values("exp_min", ascending=False).itertuples(index=False):
            _cell(ws, r, 1, game)
            _cell(ws, r, 2, team)
            _cell(ws, r, 3, p.player)
            _cell(ws, r, 4, round(p.p_play, 3), "0%", color=BLUE)
            _cell(ws, r, 5, round(p.exp_min, 1), "0.0", color=BLUE)
            _cell(ws, r, 6, round(p.p_out, 2) if p.p_out else None, "0%", color=BLUE)
            col = 7
            for st in stats:
                _cell(ws, r, col, round(getattr(p, st), 2), "0.0", color=BLUE)
                _cell(ws, r, col + 1, round(getattr(p, f"{st}_sd"), 2), "0.0", color=BLUE)
                _cell(ws, r, col + 2, getattr(p, f"{st}_line"), "0.0", color=BLUE, bold=True)
                _cell(ws, r, col + 3, round(getattr(p, f"{st}_p_over"), 3), "0.0%", color=BLUE)
                col += 4
            r += 1
        last = r - 1
        # Consistency checks as live formulas.
        _cell(ws, r, 3, f"{team} expected minutes of listed players", bold=True)
        _cell(ws, r, 5, f"=SUMPRODUCT(D{first}:D{last},E{first}:E{last})", "0.0", bold=True)
        _cell(ws, r + 1, 3, f"{team} model team score | listed players' points", bold=True)
        _cell(ws, r + 1, 5, round(float(d["team_pts"].iloc[0]), 1), "0.0", color=BLUE,
              bold=True)
        _cell(ws, r + 1, 7, f"=SUMPRODUCT(D{first}:D{last},G{first}:G{last})", "0.0", bold=True)
        r += 2
        # Players ruled out: listed for reference only, no projections, not in the totals.
        for p in ruled_out.sort_values("base_min", ascending=False).itertuples(index=False):
            for col, v in ((1, game), (2, team), (3, p.player),
                           (5, round(p.base_min, 1)), (6, round(p.p_out, 2)),
                           (7, "Ruled out (injury report): no projection")):
                c = ws.cell(row=r, column=col, value=v)
                c.font = grey
                if col == 5:
                    c.number_format = "0.0"
                if col == 6:
                    c.number_format = "0%"
            r += 1
        r += 1
    ws.freeze_panes = "D5"


def _sheet_ladder(ws, allprops: pd.DataFrame, props_model) -> int:
    """Long table of P(over/under) for half-point lines around each projection."""
    from .props import ladder_lines

    _title(ws, "Prop over/under ladder",
           "Negative binomial probabilities if the player plays. Used by the Prop checker sheet.")
    headers = ["Key", "Game", "Team", "Player", "Stat", "Line", "P(over)", "P(under)",
               "Fair over odds", "Fair under odds"]
    _header(ws, 4, headers, [30, 10, 6, 24, 6, 7, 9, 9, 10, 10])
    r = 5
    playing = allprops[allprops["p_play"] > 0] if not allprops.empty else allprops
    for p in playing.itertuples(index=False):
        for st in ("pts", "reb", "ast", "fgm3", "pra"):
            mean = getattr(p, st)
            for line in ladder_lines(mean, st):
                po, _, pu = pipeline.prop_probs(props_model, mean, st, p.exp_min, line)
                _cell(ws, r, 1, f"{p.player}|{st}|{line:g}")
                _cell(ws, r, 2, p.game)
                _cell(ws, r, 3, p.team)
                _cell(ws, r, 4, p.player)
                _cell(ws, r, 5, st)
                _cell(ws, r, 6, line, "0.0")
                _cell(ws, r, 7, round(po, 4), "0.0%", color=BLUE)
                _cell(ws, r, 8, round(pu, 4), "0.0%", color=BLUE)
                _cell(ws, r, 9, f"=IF(G{r}>0,1/G{r},\"\")", "0.00")
                _cell(ws, r, 10, f"=IF(H{r}>0,1/H{r},\"\")", "0.00")
                r += 1
    ws.freeze_panes = "B5"
    ws.auto_filter.ref = f"A4:J{r - 1}"
    return r - 1


def _sheet_checker(ws, allprops: pd.DataFrame, ladder_last: int) -> None:
    _title(ws, "Prop line checker",
           "Type a bookmaker line and odds in the yellow cells; probability and EV come from the "
           "Prop ladder. Player = name exactly as in the list (column N); stat = pts, reb, ast, "
           "fgm3 or pra; line must end in .5.")
    headers = ["Player", "Stat", "Line", "Over odds", "Under odds", "P(over)", "P(under)",
               "EV over", "EV under", "Better side"]
    _header(ws, 4, headers, [24, 7, 7, 9, 9, 9, 9, 9, 9, 10])
    keys = f"'Prop ladder'!$A$5:$A${ladder_last}"
    pover = f"'Prop ladder'!$G$5:$G${ladder_last}"
    punder = f"'Prop ladder'!$H$5:$H${ladder_last}"
    names = sorted(allprops.loc[allprops["p_play"] > 0, "player"].unique()) \
        if not allprops.empty else []
    ws.cell(row=4, column=14, value="Players").font = _f(bold=True)
    for i, n in enumerate(names, 5):
        ws.cell(row=i, column=14, value=n).font = _f()
    ws.column_dimensions["N"].width = 26
    pv = DataValidation(type="list", formula1=f"=$N$5:$N${4 + max(1, len(names))}",
                        allow_blank=True)
    sv = DataValidation(type="list", formula1='"pts,reb,ast,fgm3,pra"', allow_blank=True)
    ws.add_data_validation(pv)
    ws.add_data_validation(sv)
    example = None
    if not allprops.empty:
        top = allprops[allprops["p_play"] > 0].sort_values("pts", ascending=False).iloc[0]
        example = (top["player"], "pts", float(top["pts_line"]), 1.87, 1.93)
    for r in range(5, 25):
        vals = example if (r == 5 and example) else (None,) * 5
        for j, v in enumerate(vals, 1):
            fmt = "0.0" if j == 3 else ("0.00" if j in (4, 5) else None)
            _cell(ws, r, j, v, fmt, color=BLUE, fill=INPUT_FILL)
        pv.add(f"A{r}")
        sv.add(f"B{r}")
        key = f'A{r}&"|"&B{r}&"|"&C{r}'
        _cell(ws, r, 6, f'=IF(A{r}="","",IFERROR(INDEX({pover},MATCH({key},{keys},0)),'
                        f'"not in ladder"))', "0.0%")
        _cell(ws, r, 7, f'=IF(A{r}="","",IFERROR(INDEX({punder},MATCH({key},{keys},0)),'
                        f'"not in ladder"))', "0.0%")
        _cell(ws, r, 8, f'=IF(OR(D{r}="",NOT(ISNUMBER(F{r}))),"",F{r}*(D{r}-1)-G{r})',
              "+0.0%;-0.0%;0.0%")
        _cell(ws, r, 9, f'=IF(OR(E{r}="",NOT(ISNUMBER(G{r}))),"",G{r}*(E{r}-1)-F{r})',
              "+0.0%;-0.0%;0.0%")
        _cell(ws, r, 10, f'=IF(OR(H{r}="",I{r}=""),"",IF(MAX(H{r},I{r})<0,"no bet",'
                         f'IF(H{r}>I{r},"Over","Under")))')
    ws.cell(row=26, column=1, value=(
        "Row 5 is an example (the top projected scorer at the fair line with typical odds); overwrite it. "
        "EV = expected profit per 1 staked. Props are void if the player does not play.")
    ).font = _f(italic=True, color="595959")


def _sheet_notes(ws, model, stamp, inj) -> None:
    ws.column_dimensions["A"].width = 120
    lines = [
        ("About this file", True),
        ("Generated by elbet (EuroLeague betting model) from the repository's euroleague-model/ "
         "folder; rerun with: elbet picks --market-weight 0.5, then elbet excel --date YYYY-MM-DD.",
         False),
        (f"Odds: The Odds API (the-odds-api.com), ~20 EU bookmakers, snapshot {stamp} UTC. "
         "'Best odds' is the highest price across books; check your own bookmaker's price.", False),
        ("Results/box scores: official EuroLeague APIs (api-live.euroleague.net, "
         "live.euroleague.net), 2012-13 to date.", False),
        ("Injuries: basketballsphere.com/en/injuries/euroleague/ (updated "
         f"{inj['updated'].iloc[0] if 'updated' in inj and len(inj) else '?'}). "
         "Out = 100% absent, Doubtful 75%, Questionable/Day to day 50%.", False),
        ("", False),
        ("Model", True),
        ("Team ratings: opponent-adjusted offence/defence per 100 possessions and pace "
         "(recency-weighted ridge regression) plus Elo. Game model adds rest, venue, phase, "
         "league scoring trend and player availability.", False),
        (f"Outcome spread: margin SD {model.margin_sd:.1f} pts, total SD {model.total_sd:.1f} pts "
         "(normal, discretised to whole points; pushes handled).", False),
        ("Backtest 2015-2026 (3,375 games, walk-forward): 66.1% winners, log loss 0.610 vs 0.621 "
         "for Elo; margin MAE 9.3, total MAE 13.2 pts.", False),
        ("Picks blend model and market 50/50. Stakes: quarter Kelly, max 3% per bet.", False),
        ("", False),
        ("Caveats", True),
        ("Paper trading only until closing-line value is positive over 100+ bets. "
         "The model cannot see late scratches, domestic-league fatigue or line-ups.", False),
        ("Bet only with licensed bookmakers where it is legal for you.", False),
    ]
    for i, (t, bold) in enumerate(lines, 1):
        c = ws.cell(row=i, column=1, value=t)
        c.font = _f(bold=bold, size=12 if bold else 10)
        c.alignment = Alignment(wrap_text=True)
